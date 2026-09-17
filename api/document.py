import base64
import io
import json
import os
import zipfile

from http.server import BaseHTTPRequestHandler
from pathlib import Path

from docx import Document
from openai import OpenAI
from pypdf import PdfReader


# ============================================================
# CONFIGURACIÓN
# ============================================================

MAX_REQUEST_BYTES = 4_300_000

# Usamos un límite conservador porque Base64 incrementa
# el tamaño del archivo al enviarlo dentro de JSON.
MAX_DOCUMENT_BYTES = 2_500_000

MAX_EXTRACTED_CHARACTERS = 30_000

TRANSLATION_CHUNK_SIZE = 6_000


SUPPORTED_EXTENSIONS = {
    ".pdf",
    ".docx",
    ".txt"
}


LANGUAGES = {
    "es": "Spanish",
    "en": "English"
}


OPENAI_DOCUMENT_MODEL = os.environ.get(
    "OPENAI_DOCUMENT_MODEL",
    "gpt-4.1-mini"
)


RAW_ALLOWED_ORIGINS = os.environ.get(
    "ALLOWED_ORIGIN",
    ""
)


ALLOWED_ORIGINS = {
    origin.strip().rstrip("/")
    for origin in RAW_ALLOWED_ORIGINS.split(",")
    if origin.strip()
}


# ============================================================
# EXTRACTOR DE DOCUMENTOS
# ============================================================

class DocumentExtractor:

    def __init__(self, file_bytes, file_name):
        self.file_bytes = file_bytes
        self.file_name = file_name

        self.extension = (
            Path(file_name)
            .suffix
            .lower()
        )


    # --------------------------------------------------------
    # VALIDACIÓN
    # --------------------------------------------------------

    def validate(self):

        if self.extension not in SUPPORTED_EXTENSIONS:
            raise ValueError(
                "Formato no permitido. "
                "Usa archivos PDF, DOCX o TXT."
            )


        if not self.file_bytes:
            raise ValueError(
                "El documento está vacío."
            )


        if len(self.file_bytes) > MAX_DOCUMENT_BYTES:
            raise ValueError(
                "El documento supera el tamaño máximo permitido."
            )


        if self.extension == ".pdf":
            self._validate_pdf()


        elif self.extension == ".docx":
            self._validate_docx()


        elif self.extension == ".txt":
            self._validate_txt()


    def _validate_pdf(self):

        if not self.file_bytes.startswith(
            b"%PDF-"
        ):
            raise ValueError(
                "El archivo seleccionado no parece ser un PDF válido."
            )


    def _validate_docx(self):

        try:

            with zipfile.ZipFile(
                io.BytesIO(
                    self.file_bytes
                )
            ) as archive:

                if (
                    "word/document.xml"
                    not in archive.namelist()
                ):
                    raise ValueError(
                        "El archivo seleccionado no parece ser un DOCX válido."
                    )

        except zipfile.BadZipFile:

            raise ValueError(
                "El archivo seleccionado no parece ser un DOCX válido."
            )


    def _validate_txt(self):

        if b"\x00" in self.file_bytes:
            raise ValueError(
                "El archivo TXT contiene datos que no parecen ser texto."
            )


    # --------------------------------------------------------
    # EXTRACCIÓN
    # --------------------------------------------------------

    def extract(self):

        self.validate()


        if self.extension == ".pdf":
            text, metadata = (
                self._extract_pdf()
            )

        elif self.extension == ".docx":
            text, metadata = (
                self._extract_docx()
            )

        else:
            text, metadata = (
                self._extract_txt()
            )


        text = self._normalize_text(
            text
        )


        if not text:

            raise ValueError(
                "El documento no contiene texto procesable."
            )


        if (
            len(text) >
            MAX_EXTRACTED_CHARACTERS
        ):

            raise ValueError(
                "El documento contiene demasiado texto. "
                f"El máximo permitido actualmente es de "
                f"{MAX_EXTRACTED_CHARACTERS:,} caracteres."
            )


        metadata[
            "characters"
        ] = len(text)


        return text, metadata


    # --------------------------------------------------------
    # PDF
    # --------------------------------------------------------

    def _extract_pdf(self):

        try:

            reader = PdfReader(
                io.BytesIO(
                    self.file_bytes
                )
            )

        except Exception as error:

            raise ValueError(
                "No fue posible abrir el archivo PDF."
            ) from error


        sections = []


        for index, page in enumerate(
            reader.pages,
            start=1
        ):

            try:

                page_text = (
                    page.extract_text()
                    or ""
                ).strip()

            except Exception:

                page_text = ""


            if page_text:

                sections.append(
                    f"[Página {index}]\n"
                    f"{page_text}"
                )


        text = "\n\n".join(
            sections
        )


        metadata = {
            "type": "PDF",
            "pages": len(
                reader.pages
            )
        }


        return text, metadata


    # --------------------------------------------------------
    # DOCX
    # --------------------------------------------------------

    def _extract_docx(self):

        try:

            document = Document(
                io.BytesIO(
                    self.file_bytes
                )
            )

        except Exception as error:

            raise ValueError(
                "No fue posible abrir el documento de Word."
            ) from error


        sections = []


        for paragraph in document.paragraphs:

            text = (
                paragraph.text
                or ""
            ).strip()


            if not text:
                continue


            style_name = ""

            try:

                style_name = (
                    paragraph.style.name
                    or ""
                )

            except Exception:

                pass


            if style_name.lower().startswith(
                "heading"
            ):

                level = "##"

                parts = (
                    style_name
                    .split()
                )


                if (
                    parts and
                    parts[-1].isdigit()
                ):

                    heading_level = min(
                        int(
                            parts[-1]
                        ),
                        6
                    )

                    level = (
                        "#"
                        * heading_level
                    )


                sections.append(
                    f"{level} {text}"
                )

            else:

                sections.append(
                    text
                )


        # También extraemos tablas sencillas.
        for table_index, table in enumerate(
            document.tables,
            start=1
        ):

            table_lines = [
                f"[Tabla {table_index}]"
            ]


            for row in table.rows:

                cells = [
                    (
                        cell.text
                        or ""
                    ).strip()
                    for cell in row.cells
                ]


                if any(cells):

                    table_lines.append(
                        " | ".join(
                            cells
                        )
                    )


            if len(table_lines) > 1:

                sections.append(
                    "\n".join(
                        table_lines
                    )
                )


        text = "\n\n".join(
            sections
        )


        metadata = {
            "type": "DOCX",
            "paragraphs": len(
                document.paragraphs
            ),
            "tables": len(
                document.tables
            )
        }


        return text, metadata


    # --------------------------------------------------------
    # TXT
    # --------------------------------------------------------

    def _extract_txt(self):

        try:

            text = self.file_bytes.decode(
                "utf-8-sig"
            )

        except UnicodeDecodeError:

            try:

                text = self.file_bytes.decode(
                    "latin-1"
                )

            except UnicodeDecodeError as error:

                raise ValueError(
                    "No fue posible interpretar el archivo TXT."
                ) from error


        metadata = {
            "type": "TXT"
        }


        return text, metadata


    # --------------------------------------------------------
    # NORMALIZAR TEXTO
    # --------------------------------------------------------

    def _normalize_text(
        self,
        text
    ):

        text = (
            text
            .replace(
                "\r\n",
                "\n"
            )
            .replace(
                "\r",
                "\n"
            )
        )


        lines = []


        previous_blank = False


        for line in text.split(
            "\n"
        ):

            cleaned = line.rstrip()


            if cleaned:

                lines.append(
                    cleaned
                )

                previous_blank = False

            else:

                if not previous_blank:

                    lines.append(
                        ""
                    )

                    previous_blank = True


        return (
            "\n"
            .join(lines)
            .strip()
        )


# ============================================================
# SERVICIO DE TRADUCCIÓN
# ============================================================

class DocumentTranslationService:

    def __init__(
        self,
        client,
        model
    ):

        self.client = client

        self.model = model


    # --------------------------------------------------------
    # DIVIDIR DOCUMENTO
    # --------------------------------------------------------

    def _create_chunks(
        self,
        text
    ):

        paragraphs = text.split(
            "\n\n"
        )


        chunks = []

        current = ""


        for paragraph in paragraphs:

            paragraph = (
                paragraph.strip()
            )


            if not paragraph:
                continue


            candidate = (
                paragraph
                if not current
                else (
                    current
                    + "\n\n"
                    + paragraph
                )
            )


            if (
                len(candidate)
                <= TRANSLATION_CHUNK_SIZE
            ):

                current = candidate

                continue


            if current:

                chunks.append(
                    current
                )

                current = ""


            # Párrafo excepcionalmente largo.
            if (
                len(paragraph)
                > TRANSLATION_CHUNK_SIZE
            ):

                start = 0


                while (
                    start
                    < len(paragraph)
                ):

                    end = (
                        start
                        + TRANSLATION_CHUNK_SIZE
                    )


                    chunks.append(
                        paragraph[
                            start:end
                        ]
                    )


                    start = end

            else:

                current = paragraph


        if current:

            chunks.append(
                current
            )


        return chunks


    # --------------------------------------------------------
    # TRADUCIR
    # --------------------------------------------------------

    def translate(
        self,
        text,
        source_language,
        target_language
    ):

        source_name = (
            LANGUAGES[
                source_language
            ]
        )

        target_name = (
            LANGUAGES[
                target_language
            ]
        )


        chunks = (
            self._create_chunks(
                text
            )
        )


        if not chunks:

            raise ValueError(
                "No existe contenido para traducir."
            )


        translated_chunks = []


        total = len(
            chunks
        )


        for index, chunk in enumerate(
            chunks,
            start=1
        ):

            instructions = (
                "You are a professional document translator "
                "specialized in Spanish and English. "
                "Translate accurately and naturally while preserving "
                "the structure of the source document. "
                "Preserve titles, paragraph breaks, page markers, "
                "table markers, numbers, dates, units, acronyms, "
                "proper names and technical terminology when appropriate. "
                "Do not summarize, omit, explain, or add information. "
                "Return only the translated document fragment."
            )


            prompt = (
                f"Source language: {source_name}\n"
                f"Target language: {target_name}\n"
                f"Document fragment: {index} of {total}\n\n"
                "Translate the following content:\n\n"
                f"{chunk}"
            )


            response = (
                self.client.responses.create(
                    model=self.model,
                    instructions=instructions,
                    input=prompt,
                    store=False,
                    max_output_tokens=8000
                )
            )


            translated = (
                response.output_text
                or ""
            ).strip()


            if not translated:

                raise ValueError(
                    "La IA devolvió una sección de traducción vacía."
                )


            translated_chunks.append(
                translated
            )


        return (
            "\n\n"
            .join(
                translated_chunks
            ),
            total
        )


# ============================================================
# HANDLER VERCEL
# ============================================================

class handler(BaseHTTPRequestHandler):

    def log_message(
        self,
        format,
        *args
    ):
        return


    # --------------------------------------------------------
    # ORIGEN
    # --------------------------------------------------------

    def _get_origin(self):

        return (
            self.headers
            .get(
                "Origin",
                ""
            )
            .rstrip("/")
        )


    def _origin_is_allowed(
        self
    ):

        origin = (
            self._get_origin()
        )


        if not origin:

            return False


        return (
            origin
            in ALLOWED_ORIGINS
        )


    # --------------------------------------------------------
    # CORS
    # --------------------------------------------------------

    def _add_cors_headers(
        self
    ):

        origin = (
            self._get_origin()
        )


        if (
            origin and
            origin in ALLOWED_ORIGINS
        ):

            self.send_header(
                "Access-Control-Allow-Origin",
                origin
            )

            self.send_header(
                "Vary",
                "Origin"
            )


        self.send_header(
            "Access-Control-Allow-Methods",
            "POST, OPTIONS"
        )


        self.send_header(
            "Access-Control-Allow-Headers",
            "Content-Type"
        )


    # --------------------------------------------------------
    # JSON
    # --------------------------------------------------------

    def _send_json(
        self,
        status_code,
        payload
    ):

        data = json.dumps(
            payload,
            ensure_ascii=False
        ).encode(
            "utf-8"
        )


        self.send_response(
            status_code
        )


        self._add_cors_headers()


        self.send_header(
            "Content-Type",
            "application/json; charset=utf-8"
        )


        self.send_header(
            "Content-Length",
            str(
                len(data)
            )
        )


        self.end_headers()


        self.wfile.write(
            data
        )


    # --------------------------------------------------------
    # OPTIONS
    # --------------------------------------------------------

    def do_OPTIONS(
        self
    ):

        if not self._origin_is_allowed():

            self.send_response(
                403
            )

            self.end_headers()

            return


        self.send_response(
            204
        )


        self._add_cors_headers()


        self.end_headers()


    # --------------------------------------------------------
    # POST
    # --------------------------------------------------------

    def do_POST(
        self
    ):

        # ----------------------------------------------------
        # CORS
        # ----------------------------------------------------

        if not self._origin_is_allowed():

            self._send_json(
                403,
                {
                    "success": False,
                    "error": (
                        "Origen no autorizado."
                    )
                }
            )

            return


        # ----------------------------------------------------
        # API KEY
        # ----------------------------------------------------

        api_key = os.environ.get(
            "OPENAI_API_KEY"
        )


        if not api_key:

            self._send_json(
                500,
                {
                    "success": False,
                    "error": (
                        "El servicio no está configurado correctamente."
                    )
                }
            )

            return


        # ----------------------------------------------------
        # CONTENT TYPE
        # ----------------------------------------------------

        content_type = (
            self.headers
            .get(
                "Content-Type",
                ""
            )
            .lower()
        )


        if (
            "application/json"
            not in content_type
        ):

            self._send_json(
                415,
                {
                    "success": False,
                    "error": (
                        "El contenido debe enviarse en formato JSON."
                    )
                }
            )

            return


        # ----------------------------------------------------
        # TAMAÑO DE SOLICITUD
        # ----------------------------------------------------

        try:

            content_length = int(
                self.headers.get(
                    "Content-Length",
                    "0"
                )
            )

        except ValueError:

            self._send_json(
                400,
                {
                    "success": False,
                    "error": (
                        "Tamaño de solicitud inválido."
                    )
                }
            )

            return


        if content_length <= 0:

            self._send_json(
                400,
                {
                    "success": False,
                    "error": (
                        "La solicitud está vacía."
                    )
                }
            )

            return


        if (
            content_length
            > MAX_REQUEST_BYTES
        ):

            self._send_json(
                413,
                {
                    "success": False,
                    "error": (
                        "La solicitud supera el tamaño permitido."
                    )
                }
            )

            return


        # ----------------------------------------------------
        # LEER JSON
        # ----------------------------------------------------

        try:

            raw_body = (
                self.rfile.read(
                    content_length
                )
            )


            body = json.loads(
                raw_body.decode(
                    "utf-8"
                )
            )

        except (
            json.JSONDecodeError,
            UnicodeDecodeError
        ):

            self._send_json(
                400,
                {
                    "success": False,
                    "error": (
                        "El cuerpo de la solicitud no contiene JSON válido."
                    )
                }
            )

            return


        if not isinstance(
            body,
            dict
        ):

            self._send_json(
                400,
                {
                    "success": False,
                    "error": (
                        "Solicitud inválida."
                    )
                }
            )

            return


        # ----------------------------------------------------
        # DATOS
        # ----------------------------------------------------

        file_name = str(
            body.get(
                "file_name",
                ""
            )
        ).strip()


        file_base64 = str(
            body.get(
                "file_base64",
                ""
            )
        ).strip()


        source_language = (
            str(
                body.get(
                    "source_language",
                    ""
                )
            )
            .strip()
            .lower()
        )


        target_language = (
            str(
                body.get(
                    "target_language",
                    ""
                )
            )
            .strip()
            .lower()
        )


        # ----------------------------------------------------
        # VALIDACIONES
        # ----------------------------------------------------

        if not file_name:

            self._send_json(
                400,
                {
                    "success": False,
                    "error": (
                        "No se recibió el nombre del archivo."
                    )
                }
            )

            return


        if not file_base64:

            self._send_json(
                400,
                {
                    "success": False,
                    "error": (
                        "Selecciona un documento antes de continuar."
                    )
                }
            )

            return


        if (
            source_language
            not in LANGUAGES
        ):

            self._send_json(
                400,
                {
                    "success": False,
                    "error": (
                        "Idioma de origen no válido."
                    )
                }
            )

            return


        if (
            target_language
            not in LANGUAGES
        ):

            self._send_json(
                400,
                {
                    "success": False,
                    "error": (
                        "Idioma de destino no válido."
                    )
                }
            )

            return


        if (
            source_language
            == target_language
        ):

            self._send_json(
                400,
                {
                    "success": False,
                    "error": (
                        "Los idiomas de origen y destino deben ser diferentes."
                    )
                }
            )

            return


        # ----------------------------------------------------
        # DECODIFICAR DOCUMENTO
        # ----------------------------------------------------

        try:

            file_bytes = (
                base64.b64decode(
                    file_base64,
                    validate=True
                )
            )

        except Exception:

            self._send_json(
                400,
                {
                    "success": False,
                    "error": (
                        "El documento codificado es inválido."
                    )
                }
            )

            return


        if (
            len(file_bytes)
            > MAX_DOCUMENT_BYTES
        ):

            self._send_json(
                413,
                {
                    "success": False,
                    "error": (
                        "El documento supera el tamaño máximo permitido de 2.5 MB."
                    )
                }
            )

            return


        # ----------------------------------------------------
        # EXTRAER + TRADUCIR
        # ----------------------------------------------------

        try:

            extractor = (
                DocumentExtractor(
                    file_bytes=file_bytes,
                    file_name=file_name
                )
            )


            original_text, metadata = (
                extractor.extract()
            )


            client = OpenAI(
                api_key=api_key
            )


            translator = (
                DocumentTranslationService(
                    client=client,
                    model=OPENAI_DOCUMENT_MODEL
                )
            )


            translation, chunks = (
                translator.translate(
                    text=original_text,
                    source_language=source_language,
                    target_language=target_language
                )
            )


            metadata[
                "chunks"
            ] = chunks


            self._send_json(
                200,
                {
                    "success": True,
                    "file_name": file_name,
                    "source_language": source_language,
                    "target_language": target_language,
                    "original_text": original_text,
                    "translation": translation,
                    "metadata": metadata,
                    "message": (
                        "Documento procesado y traducido correctamente."
                    )
                }
            )


        except ValueError as error:

            self._send_json(
                422,
                {
                    "success": False,
                    "error": str(
                        error
                    )
                }
            )


        except Exception as error:

            print(
                "Document translation error:",
                type(error).__name__,
                str(error)
            )


            self._send_json(
                502,
                {
                    "success": False,
                    "error": (
                        "No fue posible procesar el documento. "
                        "Intenta nuevamente."
                    )
                }
            )